"""Text extraction from PDF, DOCX and plain-text uploads.

Extraction is intentionally conservative.  Nothing in an uploaded file is
executed or resolved: PDFs are parsed page by page with ``pypdf``, DOCX files
are read as a ZIP of XML parts by ``python-docx``, and external references,
embedded scripts and remote resources are simply never followed.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import DocumentExtractionError
from app.security.file_validation import DocumentFormat, ValidatedUpload

#: Collapse runs of whitespace but preserve paragraph breaks, which carry the
#: clause structure that chunking relies on.
_HORIZONTAL_WHITESPACE = re.compile(r"[ \t\r\f\v]+")
_EXCESS_NEWLINES = re.compile(r"\n{3,}")
_MIN_USEFUL_CHARACTERS = 40


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    """Normalised text plus the provenance needed for citations."""

    text: str
    page_count: int
    character_count: int
    truncated: bool


def normalise_whitespace(raw: str) -> str:
    """Collapse whitespace while preserving paragraph boundaries.

    Args:
        raw: Text as produced by a format-specific extractor.

    Returns:
        Text with horizontal whitespace collapsed and blank runs limited.
    """
    collapsed = _HORIZONTAL_WHITESPACE.sub(" ", raw)
    collapsed = "\n".join(line.strip() for line in collapsed.split("\n"))
    return _EXCESS_NEWLINES.sub("\n\n", collapsed).strip()


def _extract_pdf(content: bytes) -> tuple[str, int]:
    """Extract text from a PDF payload."""
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentExtractionError(
                "This PDF is password protected. Please upload an unlocked copy."
            )
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentExtractionError:
        raise
    except (PdfReadError, ValueError, OSError) as exc:
        raise DocumentExtractionError(
            "This PDF could not be read. It may be corrupted.",
            detail=str(exc),
        ) from exc
    return "\n\n".join(pages), len(pages)


def _extract_docx(content: bytes) -> tuple[str, int]:
    """Extract paragraph and table text from a DOCX payload."""
    try:
        document = docx.Document(io.BytesIO(content))
    except (ValueError, KeyError, OSError) as exc:
        raise DocumentExtractionError(
            "This Word document could not be read. It may be corrupted.",
            detail=str(exc),
        ) from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks), 1


def _extract_text(content: bytes) -> tuple[str, int]:
    """Decode a plain-text payload as UTF-8."""
    return content.decode("utf-8", errors="replace"), 1


def extract_document(upload: ValidatedUpload, *, max_characters: int) -> ExtractedDocument:
    """Extract normalised text from a validated upload.

    Args:
        upload: An upload that has already passed format and size validation.
        max_characters: Hard ceiling on retained characters; longer documents
            are truncated so a single file cannot exhaust the request budget.

    Returns:
        The extracted document.

    Raises:
        DocumentExtractionError: If the file yields too little usable text,
            which in practice means a scanned image-only PDF.
    """
    extractors = {
        DocumentFormat.PDF: _extract_pdf,
        DocumentFormat.DOCX: _extract_docx,
        DocumentFormat.TEXT: _extract_text,
    }
    raw, page_count = extractors[upload.document_format](upload.content)
    text = normalise_whitespace(raw)

    if len(text) < _MIN_USEFUL_CHARACTERS:
        raise DocumentExtractionError(
            "No readable text was found in this file. If it is a scan or a photo, "
            "please upload a text-based copy instead.",
            detail=f"Extracted {len(text)} characters.",
        )

    truncated = len(text) > max_characters
    if truncated:
        text = text[:max_characters]

    return ExtractedDocument(
        text=text,
        page_count=page_count,
        character_count=len(text),
        truncated=truncated,
    )
